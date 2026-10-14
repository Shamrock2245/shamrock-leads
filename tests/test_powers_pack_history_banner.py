"""Powers-pack transfer banner is per range, not global.

Mocked Mongo only. Power numbers are TEST- fixtures.
"""

from __future__ import annotations

import asyncio
from io import BytesIO
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from dashboard.services.powers_pack import (
    TRANSFER_HISTORY_BANNER,
    assemble_powers_pack,
    build_powers_pack_xlsx,
)


class _Col:
    def __init__(self, docs):
        self.docs = list(docs)

    def find(self, query=None, projection=None):
        return self

    async def to_list(self, limit):
        return [dict(d) for d in self.docs[:limit]]


def _db(*, inventory, audits=None, bonds=None):
    return {
        "poa_inventory": _Col(inventory),
        "active_bonds": _Col(bonds or []),
        "audit_events": _Col(audits or []),
    }


def _post_172_inventory():
    return [
        {
            "surety_id": "osi",
            "poa_number": "TEST-2001",
            "status": "assigned",
            "bond_case_id": "TEST-CASE-B",
            "reassigned_from": [{
                "holder": "Prior Holder",
                "case": "TEST-CASE-A",
                "timestamp": "2026-09-12T15:00:00+00:00",
                "actor": "office@example.invalid",
            }],
        },
        {
            "surety_id": "osi",
            "poa_number": "TEST-2002",
            "status": "available",
            "release_history": [{
                "released_at": "2026-09-04T00:00:00+00:00",
                "released_by": "first@example.invalid",
                "release_reason": "bond_renewal",
            }],
            "released_at": "2026-09-18T00:00:00+00:00",
            "released_by": "office@example.invalid",
            "release_reason": "manual",
        },
    ]


def _post_172_audits():
    return [{
        "entity_type": "poa",
        "entity_id": "TEST-2001",
        "action": "poa_reassigned",
        "timestamp": "2026-09-12T15:00:00+00:00",
        "actor": "office@example.invalid",
        "details": {
            "from_case": "TEST-CASE-A",
            "to_case": "TEST-CASE-B",
            "poa_number": "TEST-2001",
        },
    }]


def _legacy_string_power():
    return {
        "surety_id": "osi",
        "poa_number": "TEST-2099",
        "status": "assigned",
        "reassigned_from": "TEST-CASE-OLD",
        "bond_case_id": "TEST-CASE-NEW",
        "used_at": "2026-09-20T12:00:00+00:00",
    }


def _pack(db):
    return asyncio.run(assemble_powers_pack(
        db,
        surety_id="osi",
        pack="transfer",
        start_date="2026-09-01",
        end_date="2026-09-30",
    ))


def _workbook_text(xlsx: bytes) -> str:
    wb = load_workbook(BytesIO(xlsx))
    parts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    parts.append(str(cell.value))
    return "\n".join(parts)


def _cover_map(xlsx: bytes) -> dict:
    wb = load_workbook(BytesIO(xlsx))
    ws = wb["OSI Cover"]
    found = {}
    for row in ws.iter_rows(min_col=1, max_col=2, values_only=True):
        if row[0]:
            found[row[0]] = row[1]
    return found


def _assert_no_private_keys(rows):
    for row in rows:
        assert "_string_reassign" not in row
        assert "_legacy" not in row


def test_post_172_range_omits_banner_and_marks_history_complete():
    pack = _pack(_db(inventory=_post_172_inventory(), audits=_post_172_audits()))
    rows = pack["transfer"]["rows"]
    reassigns = [row for row in rows if row["event_type"] == "reassign"]
    releases = [row for row in rows if row["event_type"] == "release"]

    assert [(row["poa_number"], row["from_ref"], row["event_date"]) for row in reassigns] == [
        ("TEST-2001", "TEST-CASE-A", "2026-09-12"),
    ]
    assert reassigns[0]["evidence_complete"] is True
    assert "audit_events.poa_reassigned" in reassigns[0]["sources"]
    assert "poa_inventory.reassigned_from" in reassigns[0]["sources"]

    assert [(row["poa_number"], row["event_date"], row["sources"]) for row in releases] == [
        ("TEST-2002", "2026-09-04", ["poa_inventory.release_history"]),
        ("TEST-2002", "2026-09-18", ["poa_inventory.released_at"]),
    ]
    assert [row["evidence_complete"] for row in releases] == [True, True]
    _assert_no_private_keys(rows)

    assert pack["history_complete"] is True
    assert pack["transfer"]["history_complete"] is True
    assert pack["transfer_banner"] is None
    assert pack["transfer"]["banner"] is None

    raw = build_powers_pack_xlsx(pack, pack="transfer")
    text = _workbook_text(raw)
    assert TRANSFER_HISTORY_BANNER not in text
    assert "best-effort" not in text
    cover = _cover_map(raw)
    assert cover["Transfer history complete"] == "yes"
    assert "Transfer log" not in cover
    transfer_note = load_workbook(BytesIO(raw))["OSI Transfer"]["A2"].value
    assert not transfer_note


def test_legacy_string_reassign_restores_banner():
    inventory = _post_172_inventory() + [_legacy_string_power()]
    pack = _pack(_db(inventory=inventory, audits=_post_172_audits()))
    rows = pack["transfer"]["rows"]
    by_power = {}
    for row in rows:
        by_power.setdefault(row["poa_number"], []).append(row)

    assert by_power["TEST-2001"][0]["evidence_complete"] is True
    assert [row["evidence_complete"] for row in by_power["TEST-2002"]] == [True, True]
    legacy = by_power["TEST-2099"]
    assert len(legacy) == 1
    assert legacy[0]["event_type"] == "reassign"
    assert legacy[0]["from_ref"] == "TEST-CASE-OLD"
    assert legacy[0]["event_date"] == "2026-09-20"
    assert legacy[0]["sources"] == ["poa_inventory.reassigned_from"]
    assert legacy[0]["evidence_complete"] is False
    _assert_no_private_keys(rows)

    assert pack["history_complete"] is False
    assert pack["transfer"]["history_complete"] is False
    assert pack["transfer_banner"] == TRANSFER_HISTORY_BANNER
    assert pack["transfer"]["banner"] == TRANSFER_HISTORY_BANNER

    raw = build_powers_pack_xlsx(pack, pack="transfer")
    text = _workbook_text(raw)
    assert TRANSFER_HISTORY_BANNER in text
    assert "best-effort" in text
    cover = _cover_map(raw)
    assert cover["Transfer history complete"] == "no"
    assert cover["Transfer log"] == TRANSFER_HISTORY_BANNER
    assert load_workbook(BytesIO(raw))["OSI Transfer"]["A2"].value == TRANSFER_HISTORY_BANNER


def test_unstamped_release_is_legacy():
    db = _db(
        inventory=[{
            "surety_id": "osi",
            "poa_number": "TEST-2004",
            "status": "available",
        }],
        audits=[{
            "entity_type": "poa",
            "entity_id": "TEST-2004",
            "action": "auto_released",
            "timestamp": "2026-09-08T12:00:00+00:00",
            "details": {"reason": "exonerated"},
        }],
    )
    pack = _pack(db)
    rows = pack["transfer"]["rows"]
    assert len(rows) == 1
    assert rows[0]["event_type"] == "auto_released"
    assert rows[0]["poa_number"] == "TEST-2004"
    assert rows[0]["sources"] == ["audit_events.auto_released"]
    assert rows[0]["evidence_complete"] is False
    assert pack["history_complete"] is False
    assert pack["transfer_banner"] == TRANSFER_HISTORY_BANNER
    raw = build_powers_pack_xlsx(pack, pack="transfer")
    assert TRANSFER_HISTORY_BANNER in _workbook_text(raw)
    assert _cover_map(raw)["Transfer history complete"] == "no"


@patch("dashboard.routers.reports.get_db")
def test_download_headers_follow_history_complete(mock_get_db):
    from dashboard.routers.reports import reports_bp

    app = FastAPI()
    app.include_router(reports_bp)
    client = TestClient(app)
    params = {"surety": "osi", "start_date": "2026-09-01", "end_date": "2026-09-30"}

    mock_get_db.return_value = _db(inventory=_post_172_inventory(), audits=_post_172_audits())
    complete = client.get("/api/reports/powers-pack/transfer", params=params)
    assert complete.status_code == 200
    assert complete.headers["x-history-complete"] == "true"
    assert complete.headers["x-transfer-partial"] == "false"
    assert TRANSFER_HISTORY_BANNER not in _workbook_text(complete.content)

    mock_get_db.return_value = _db(
        inventory=_post_172_inventory() + [_legacy_string_power()],
        audits=_post_172_audits(),
    )
    partial = client.get("/api/reports/powers-pack/transfer", params=params)
    assert partial.status_code == 200
    assert partial.headers["x-history-complete"] == "false"
    assert partial.headers["x-transfer-partial"] == "true"
    assert TRANSFER_HISTORY_BANNER in _workbook_text(partial.content)
