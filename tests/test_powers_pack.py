"""Carrier powers packs: contract, fail-closed surety, empty ranges, transfer evidence."""

from __future__ import annotations

import asyncio
from datetime import datetime, timezone
from io import BytesIO
from unittest.mock import patch

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from dashboard.services.powers_pack import (
    TRANSFER_HISTORY_BANNER,
    assemble_powers_pack,
    voided_powers_mongo_query,
)


class _Col:
    def __init__(self, docs):
        self.docs = list(docs)

    def find(self, query=None, projection=None):
        return self

    async def to_list(self, limit):
        return [dict(d) for d in self.docs[:limit]]


def _db(*, inventory=None, bonds=None, audits=None):
    return {
        "poa_inventory": _Col(inventory or []),
        "active_bonds": _Col(bonds or []),
        "audit_events": _Col(audits or []),
    }


def _pack(db, **kwargs):
    params = {"surety_id": "palmetto", "pack": "combined", "start_date": "2026-09-01", "end_date": "2026-09-30"}
    params.update(kwargs)
    return asyncio.run(assemble_powers_pack(db, **params))


@pytest.fixture
def reports_app():
    from dashboard.routers.reports import reports_bp

    app = FastAPI()
    app.include_router(reports_bp)
    return app


def _sheet_values(xlsx: bytes, title: str) -> list[tuple]:
    wb = load_workbook(BytesIO(xlsx))
    ws = wb[title]
    return [tuple(cell.value for cell in row) for row in ws.iter_rows()]


def _column_after_header(xlsx: bytes, title: str, header: str) -> list:
    rows = _sheet_values(xlsx, title)
    header_idx = next(i for i, row in enumerate(rows) if header in row)
    col = rows[header_idx].index(header)
    return [row[col] for row in rows[header_idx + 1:] if row[col] not in (None, "")]


def test_voided_powers_query_matches_existing_filter():
    query, _warnings = voided_powers_mongo_query("palmetto", "2026-09-01", "2026-09-30")
    assert query["status"] == "voided"
    assert query["surety_id"] == "palmetto"
    assert query["voided_at"] == {"$gte": "2026-09-01", "$lt": "2026-10-01"}


def test_empty_range_has_no_rows_and_keeps_transfer_banner():
    db = _db(
        inventory=[{
            "surety_id": "palmetto",
            "poa_number": "PSC-OUT",
            "status": "assigned",
            "date_executed": "2024-01-02",
            "bond_amount": 5000,
        }],
        bonds=[{
            "surety_id": "palmetto",
            "poa_number": "PSC-OUT",
            "status": "active",
            "bond_date": "2024-01-02",
            "bond_amount": 5000,
        }],
    )
    pack = _pack(db, start_date="2020-01-01", end_date="2020-01-31")
    assert pack["execution"]["count"] == 0
    assert pack["void"]["count"] == 0
    assert pack["transfer"]["count"] == 0
    assert pack["liability"]["count"] == 0
    assert pack["execution"]["rows"] == []
    assert pack["transfer"]["rows"] == []
    assert pack["history_complete"] is False
    assert pack["transfer"]["history_complete"] is False
    assert "best-effort" in pack["transfer_banner"]
    assert pack["transfer_banner"] == TRANSFER_HISTORY_BANNER
    assert pack["filename"] == "Palmetto_Combined_Powers_2020-01-31.xlsx"
    blob = str(pack["execution"]["rows"]) + str(pack["void"]["rows"]) + str(pack["transfer"]["rows"])
    assert "PSC-OUT" not in blob


def test_execution_uses_stored_amount_and_does_not_compute_premium():
    db = _db(inventory=[{
        "surety_id": "osi",
        "poa_number": "OSI-100",
        "poa_prefix": "OSI-P6",
        "status": "assigned",
        "date_executed": "2026-09-10",
        "bond_amount": 5000,
        "defendant_name": "Ada Lovelace",
    }])
    pack = _pack(db, surety_id="osi", pack="execution")
    assert pack["execution"]["count"] == 1
    row = pack["execution"]["rows"][0]
    assert row["poa_number"] == "OSI-100"
    assert row["bond_amount"] == 5000
    assert row["gross_premium"] is None
    assert row["surety_owed"] is None
    assert row["buf_owed"] is None
    assert row["defendant"] == "Ada Lovelace"
    assert pack["filename"].startswith("OSI_Execution_Powers_")


def test_reassign_timestamp_is_transfer_not_execution():
    db = _db(inventory=[{
        "surety_id": "palmetto",
        "poa_number": "PSC-9",
        "status": "assigned",
        "date_executed": "2026-08-01",
        "used_at": "2026-09-12T15:00:00+00:00",
        "reassigned_from": "CASE-OLD",
        "bond_case_id": "CASE-NEW",
        "bond_amount": 2500,
    }])
    pack = _pack(db)
    assert [r["poa_number"] for r in pack["execution"]["rows"]] == []
    assert pack["transfer"]["count"] == 1
    row = pack["transfer"]["rows"][0]
    assert row["event_type"] == "reassign"
    assert row["event_date"] == "2026-09-12"
    assert row["from_ref"] == "CASE-OLD"
    assert row["to_ref"] == "CASE-NEW"
    assert row["sources"] == ["poa_inventory.reassigned_from"]
    assert "bond_amount" not in row


def test_assigned_agent_alone_is_not_a_transfer():
    db = _db(inventory=[{
        "surety_id": "palmetto",
        "poa_number": "PSC-AGENT",
        "status": "assigned",
        "assigned_to_agent": "Jason Taylor",
        "date_executed": "2026-09-04",
        "bond_amount": 1000,
    }])
    pack = _pack(db)
    assert pack["transfer"]["rows"] == []
    assert pack["execution"]["count"] == 1


def test_undated_reassign_omitted_from_dated_range():
    db = _db(inventory=[{
        "surety_id": "osi",
        "poa_number": "OSI-NODATE",
        "status": "assigned",
        "reassigned_from": "CASE-1",
        "bond_case_id": "CASE-2",
    }])
    pack = _pack(db, surety_id="osi", pack="transfer")
    assert pack["transfer"]["rows"] == []
    assert pack["transfer"]["omitted_undated"] == 1
    assert any("omitted" in w for w in pack["warnings"])


def test_void_rows_match_voided_powers_filter():
    inventory = [
        {
            "surety_id": "palmetto",
            "poa_number": "PSC-V1",
            "status": "voided",
            "voided_at": "2026-09-05T18:00:00",
            "void_reason": "Wrote wrong power",
            "bond_amount": 3000,
        },
        {
            "surety_id": "palmetto",
            "poa_number": "PSC-VOUT",
            "status": "voided",
            "voided_at": "2026-08-01",
            "void_reason": "Old",
        },
        {
            "surety_id": "osi",
            "poa_number": "OSI-V",
            "status": "voided",
            "voided_at": "2026-09-05",
            "void_reason": "Other carrier",
        },
        {
            "surety_id": "palmetto",
            "poa_number": "PSC-NOTVOID",
            "status": "assigned",
            "date_executed": "2026-09-05",
        },
    ]
    pack = _pack(_db(inventory=inventory), pack="void")
    assert [r["poa_number"] for r in pack["void"]["rows"]] == ["PSC-V1"]
    assert pack["void"]["rows"][0]["void_reason"] == "Wrote wrong power"
    assert pack["void"]["rows"][0]["voided_by"] is None
    assert pack["void"]["rows"][0]["bond_amount"] == 3000


def test_release_and_renewal_merge_to_one_row():
    db = _db(
        inventory=[{
            "surety_id": "palmetto",
            "poa_number": "PSC-1",
            "status": "available",
            "released_at": datetime(2026, 9, 15, tzinfo=timezone.utc),
            "released_reason": "bond_renewal",
        }],
        bonds=[{
            "surety_id": "palmetto",
            "poa_number": "PSC-2",
            "previous_poa_number": "PSC-1",
            "last_renewed_at": "2026-09-15T16:00:00",
            "renewal_reason": "continuance",
            "bond_date": "2026-09-01",
            "status": "active",
            "defendant_name": "Grace Hopper",
            "bond_amount": 4000,
        }],
        audits=[{
            "entity_type": "poa",
            "entity_id": "PSC-1",
            "action": "auto_released",
            "timestamp": "2026-09-15T16:05:00",
            "details": {"reason": "bond_renewal"},
            "actor": "system",
        }],
    )
    pack = _pack(db)
    assert pack["transfer"]["count"] == 1
    row = pack["transfer"]["rows"][0]
    assert row["event_type"] == "renewal_poa_change"
    assert row["poa_number"] == "PSC-1"
    assert row["from_ref"] == "PSC-1"
    assert row["to_ref"] == "PSC-2"
    assert "poa_inventory.released_at" in row["sources"]
    assert "active_bonds.previous_poa_number" in row["sources"]
    assert "audit_events.auto_released" in row["sources"]
    assert pack["transfer"]["unattributed_audit_events"] == 0


def test_other_carrier_and_unknown_audit_are_not_invented_rows():
    db = _db(
        inventory=[{
            "surety_id": "osi",
            "poa_number": "OSI-1",
            "status": "available",
            "released_at": "2026-09-02",
            "release_reason": "exonerated",
        }],
        audits=[
            {
                "entity_type": "poa",
                "entity_id": "OSI-1",
                "action": "auto_released",
                "timestamp": "2026-09-02T00:00:00",
                "details": {"reason": "exonerated"},
            },
            {
                "entity_type": "poa",
                "entity_id": "UNKNOWN-9",
                "action": "auto_released",
                "timestamp": "2026-09-03T00:00:00",
                "details": {"reason": "orphan"},
            },
        ],
    )
    pack = _pack(db, surety_id="palmetto")
    assert pack["transfer"]["rows"] == []
    assert pack["transfer"]["unattributed_audit_events"] == 1


def test_manual_release_without_timestamp_is_absent():
    db = _db(inventory=[{
        "surety_id": "osi",
        "poa_number": "OSI-REL",
        "status": "available",
        "bond_case_id": None,
        "used_at": None,
    }])
    pack = _pack(db, surety_id="osi")
    assert pack["transfer"]["rows"] == []


def test_combined_workbook_sheets_and_empty_message():
    from dashboard.services.powers_pack import build_powers_pack_xlsx

    pack = _pack(_db(), start_date="2020-03-01", end_date="2020-03-31")
    raw = build_powers_pack_xlsx(pack, pack="combined")
    wb = load_workbook(BytesIO(raw))
    assert wb.sheetnames == [
        "Palmetto Cover",
        "Palmetto Execution",
        "Palmetto Void",
        "Palmetto Transfer",
        "Palmetto Liability",
    ]
    exec_text = " ".join(str(v) for row in _sheet_values(raw, "Palmetto Execution") for v in row if v)
    assert "No powers executed in this range." in exec_text
    assert _column_after_header(raw, "Palmetto Execution", "Power #") == []
    assert _column_after_header(raw, "Palmetto Transfer", "Power #") == []
    cover = " ".join(str(v) for row in _sheet_values(raw, "Palmetto Cover") for v in row if v)
    assert "best-effort" in cover
    assert "history complete" in cover.lower()


def test_execution_workbook_does_not_include_other_pack_rows():
    from dashboard.services.powers_pack import build_powers_pack_xlsx

    db = _db(inventory=[
        {
            "surety_id": "osi",
            "poa_number": "OSI-EXEC",
            "status": "assigned",
            "date_executed": "2026-09-08",
            "bond_amount": 1500,
        },
        {
            "surety_id": "osi",
            "poa_number": "OSI-VOID",
            "status": "voided",
            "voided_at": "2026-09-09",
            "void_reason": "error",
        },
    ])
    pack = _pack(db, surety_id="osi", pack="execution")
    raw = build_powers_pack_xlsx(pack, pack="execution")
    wb = load_workbook(BytesIO(raw))
    assert wb.sheetnames == ["OSI Cover", "OSI Execution"]
    powers = _column_after_header(raw, "OSI Execution", "Power #")
    assert powers == ["OSI-EXEC"]
    premiums = _column_after_header(raw, "OSI Execution", "Gross premium")
    assert premiums == []


@patch("dashboard.routers.reports.get_db")
def test_unknown_surety_is_400(mock_get_db, reports_app):
    client = TestClient(reports_app)
    resp = client.get("/api/reports/powers-pack/combined", params={"surety": "acme"})
    assert resp.status_code == 400
    body = resp.json()
    assert body["success"] is False
    assert body["error"] == "unsupported_surety"
    mock_get_db.assert_not_called()


@patch("dashboard.routers.reports.get_db")
def test_missing_surety_is_required(mock_get_db, reports_app):
    client = TestClient(reports_app)
    resp = client.get("/api/reports/powers-pack/void")
    assert resp.status_code == 400
    assert resp.json()["error"] == "surety_required"
    mock_get_db.assert_not_called()


@patch("dashboard.routers.reports.get_db")
def test_inactive_surety_fails_closed(mock_get_db, reports_app):
    client = TestClient(reports_app)
    resp = client.get("/api/reports/powers-pack/execution", params={"surety": "lexington"})
    assert resp.status_code == 400
    assert resp.json()["error"] == "surety_inactive"
    mock_get_db.assert_not_called()


@patch("dashboard.routers.reports.get_db")
def test_invalid_pack_is_400(mock_get_db, reports_app):
    client = TestClient(reports_app)
    resp = client.get("/api/reports/powers-pack/bordereau", params={"surety": "osi"})
    assert resp.status_code == 400
    assert resp.json()["error"] == "invalid_pack"
    mock_get_db.assert_not_called()


@patch("dashboard.routers.reports.get_db")
def test_endpoint_json_contract_and_xlsx_download(mock_get_db, reports_app):
    mock_get_db.return_value = _db(inventory=[{
        "surety_id": "palmetto",
        "poa_number": "PSC-LIVE",
        "status": "assigned",
        "date_executed": "2026-09-20",
        "bond_amount": 8000,
        "gross_premium": 800,
    }])
    client = TestClient(reports_app)
    js = client.get(
        "/api/reports/powers-pack/combined",
        params={"surety": "Palmetto", "start_date": "2026-09-01", "end_date": "2026-09-30", "fmt": "json"},
    )
    assert js.status_code == 200
    data = js.json()
    assert data["surety_id"] == "palmetto"
    assert data["carrier"] == "Palmetto"
    assert data["filename"] == "Palmetto_Combined_Powers_2026-09-30.xlsx"
    assert data["execution"]["count"] == 1
    assert data["execution"]["rows"][0]["gross_premium"] == 800
    assert data["void"]["count"] == 0
    assert data["history_complete"] is False

    xl = client.get(
        "/api/reports/powers-pack/combined",
        params={"surety": "PALMETTO", "start_date": "2026-09-01", "end_date": "2026-09-30"},
    )
    assert xl.status_code == 200
    assert "spreadsheetml" in xl.headers["content-type"]
    assert 'filename="Palmetto_Combined_Powers_2026-09-30.xlsx"' in xl.headers["content-disposition"]
    assert xl.headers["x-transfer-partial"] == "true"
    assert xl.headers["x-history-complete"] == "false"
    assert xl.headers["x-execution-count"] == "1"
    assert xl.content[:2] == b"PK"
    assert _column_after_header(xl.content, "Palmetto Execution", "Power #") == ["PSC-LIVE"]
