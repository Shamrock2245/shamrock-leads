"""Carrier-named OSI / Palmetto powers packs (BailSafe P0 slice C1).

Packs reuse stored POA and bond fields. They do not calculate premiums and
they do not invent transfer history.

Transfer evidence that actually exists
--------------------------------------
* ``POST /api/poa/reassign`` writes ``poa_inventory.reassigned_from`` (previous
  ``bond_case_id`` only) and overwrites ``used_at``. It does not write
  ``audit_events``, ``reassigned_at``, or from/to agent or office. The field
  is last-write: an earlier reassignment is gone.
* ``auto_release_poa`` writes ``released_at`` (datetime), ``release_reason``,
  and ``audit_events`` action ``auto_released`` (no surety on the audit row;
  surety comes from the inventory document).
* Bond renewal writes ``released_at`` / ``released_reason`` on the old power
  and ``previous_poa_number`` plus ``last_renewed_at`` on the bond.
* ``POST /api/poa/release`` sets the power back to available and does not
  store a timestamp, so those releases cannot appear.

Void rows use the same filter as ``GET /api/reports/voided-powers``:
``poa_inventory.status == voided``, exact ``surety_id``, ``voided_at`` in range.
"""
from __future__ import annotations

import io
import logging
from datetime import datetime, timezone
from typing import Any, Iterable, Optional

from dashboard.services.bond_report_xlsx import (
    AGENCY_NAME,
    REPORT_ROW_LIMIT,
    _fmt_date as parse_report_datetime,
    mongo_bond_date_filter,
    parse_report_date_window,
)
from dashboard.services.surety_registry import SURETY_REGISTRY

logger = logging.getLogger(__name__)

PACK_KINDS = ("execution", "void", "transfer", "combined")

TRANSFER_AUDIT_ACTIONS = ("auto_released", "reassigned", "poa_reassigned", "released")

# Closed book. Empty / active / monitoring stay on the liability sheet.
_CLOSED_BOND_STATUSES = frozenset({
    "void", "voided", "expired", "exonerated", "surrendered",
    "discharged", "forfeited", "closed", "cancelled",
})
_ASSIGNED_STATUSES = frozenset({"assigned", "used", "executed"})

TRANSFER_HISTORY_BANNER = (
    "Transfer log is best-effort and incomplete. "
    "Each row is backed by stored evidence only. "
    "Case reassignment: poa_inventory.reassigned_from holds the previous bond_case_id "
    "and is overwritten on the next reassign, so only the latest case move survives. "
    "POST /api/poa/reassign does not write audit_events or reassigned_at. "
    "used_at is overwritten at reassign time and is the event time when reassigned_from is set. "
    "Releases: poa_inventory.released_at plus release_reason or released_reason "
    "(auto-release and bond-renewal power swap). "
    "audit_events with entity_type poa and action auto_released are included when the power's surety_id matches this carrier. "
    "active_bonds.previous_poa_number records a renewal that replaced a power. "
    "Manual POST /api/poa/release returns a power to available and does not store a timestamp, so those releases do not appear. "
    "No agent-to-agent or office-to-office history is stored. "
    "An empty sheet means no dated evidence fell in this range."
)

EXECUTION_PACK_NOTE = (
    "Execution rows are powers with a stored execution or assign date in range: "
    "poa_inventory date_executed or executed_at, otherwise used_at when the power is "
    "assigned and has not been reassigned, plus active_bonds with a power number and bond_date in range. "
    "Amounts are stored fields only. A blank amount means that field was not on the record."
)

VOID_PACK_NOTE = (
    "Void rows use the Voided Powers filter: poa_inventory status voided, "
    "this surety_id, and voided_at inside the date range. "
    "Bond documents marked void without an inventory void are not listed."
)

LIABILITY_PACK_NOTE = (
    "Open-book bonds in range from active_bonds. Bond amount is the stored value. "
    "This pack does not recalculate premium, BUF, or surety splits. "
    "Monthly bordereau remains GET /api/reports/bordereau."
)


def powers_pack_filename(surety_id: str, pack: str, on_date: str) -> str:
    """Carrier-named download, e.g. ``Palmetto_Combined_Powers_2026-09-30.xlsx``."""
    short = SURETY_REGISTRY[surety_id]["short"]
    label = {
        "execution": "Execution",
        "void": "Void",
        "transfer": "Transfer",
        "combined": "Combined",
    }[pack]
    return f"{short}_{label}_Powers_{on_date}.xlsx"


def sheet_title(surety_id: str, name: str) -> str:
    short = SURETY_REGISTRY[surety_id]["short"]
    return f"{short} {name}"[:31]


def voided_powers_mongo_query(
    surety_id: str,
    start_date: str | None,
    end_date: str | None,
) -> tuple[dict, list[str]]:
    """Same filter as ``GET /api/reports/voided-powers`` once surety is required."""
    query: dict[str, Any] = {"status": "voided", "surety_id": surety_id}
    filt, warnings = mongo_bond_date_filter(start_date, end_date, field="voided_at")
    query.update(filt)
    return query, warnings


def _iso_day(val: Any) -> Optional[str]:
    dt = parse_report_datetime(val)
    if dt is None:
        return None
    return dt.strftime("%Y-%m-%d")


def _in_window(val: Any, start_dt: Optional[datetime], end_dt: Optional[datetime]) -> bool:
    """Calendar-day window. Matches string bounds used by voided-powers."""
    if start_dt is None and end_dt is None:
        return True
    day = _iso_day(val)
    if day is None:
        return False
    if start_dt and day < start_dt.strftime("%Y-%m-%d"):
        return False
    if end_dt and day > end_dt.strftime("%Y-%m-%d"):
        return False
    return True


def _window_active(start_dt: Optional[datetime], end_dt: Optional[datetime]) -> bool:
    return start_dt is not None or end_dt is not None


def _norm_poa(val: Any) -> str:
    return str(val or "").strip()


def _clean_text(val: Any) -> Optional[str]:
    text = str(val or "").strip()
    return text or None


def _stored_amount(doc: Optional[dict], *keys: str) -> Optional[float]:
    if not doc:
        return None
    for key in keys:
        if key not in doc:
            continue
        val = doc.get(key)
        if val is None or val == "":
            continue
        try:
            return round(float(val), 2)
        except (TypeError, ValueError):
            continue
    return None


def _first_stored(primary: Optional[dict], secondary: Optional[dict], *keys: str) -> Any:
    for doc in (primary, secondary):
        if not doc:
            continue
        for key in keys:
            if key not in doc:
                continue
            val = doc.get(key)
            if val is None or val == "":
                continue
            return val
    return None


def _stored_defendant(doc: Optional[dict]) -> Optional[str]:
    if not doc:
        return None
    name = _clean_text(doc.get("defendant_name") or doc.get("full_name"))
    if name:
        return name
    first = _clean_text(doc.get("defendant_first_name") or doc.get("first_name"))
    last = _clean_text(doc.get("defendant_last_name") or doc.get("last_name"))
    joined = " ".join(p for p in (first, last) if p)
    return joined or None


def inventory_surety_matches(doc: dict, surety_id: str) -> bool:
    return str(doc.get("surety_id") or "").strip().lower() == surety_id


def bond_matches_surety(doc: dict, surety_id: str) -> bool:
    """Exact carrier match. Unknown labels do not fall through to OSI."""
    if inventory_surety_matches(doc, surety_id):
        return True
    for key in ("surety", "insurance_company"):
        raw = str(doc.get(key) or "").strip()
        if not raw:
            continue
        if raw.strip().lower() == surety_id:
            return True
        up = raw.upper()
        if surety_id == "osi" and (up == "OSI" or "SHAUGHNAHILL" in up):
            return True
        if surety_id == "palmetto" and (up in ("PALMETTO", "PSC") or up.startswith("PALMETTO")):
            return True
    return False


def _inventory_query(surety_id: str) -> dict:
    short = SURETY_REGISTRY[surety_id]["short"]
    return {"surety_id": {"$in": [surety_id, surety_id.upper(), short, short.upper()]}}


def _bond_query(surety_id: str) -> dict:
    short = SURETY_REGISTRY[surety_id]["short"]
    labels = [surety_id, surety_id.upper(), short, short.upper()]
    if surety_id == "osi":
        labels.append("O'Shaughnahill Surety & Insurance")
    if surety_id == "palmetto":
        labels.extend(["Palmetto Surety Corporation", "PSC"])
    # De-dupe while keeping order
    seen = []
    for label in labels:
        if label not in seen:
            seen.append(label)
    return {"$or": [
        {"surety_id": {"$in": seen}},
        {"surety": {"$in": seen}},
        {"insurance_company": {"$in": seen}},
    ]}


async def _fetch(db, name: str, query: dict) -> tuple[list[dict], bool]:
    col = db[name]
    docs = await col.find(query, {"_id": 0}).to_list(REPORT_ROW_LIMIT)
    rows = [d for d in docs if isinstance(d, dict)]
    return rows, len(rows) >= REPORT_ROW_LIMIT


def _execution_stamp(inv: dict) -> tuple[Any, str] | None:
    """Return (raw date, basis) for a power execution. None when not evidenced."""
    if inv.get("date_executed") not in (None, "") or inv.get("executed_at") not in (None, ""):
        raw = inv.get("date_executed") if inv.get("date_executed") not in (None, "") else inv.get("executed_at")
        return raw, "poa_inventory.date_executed"
    # Reassign overwrites used_at. That stamp is a transfer time, not an execution.
    if _clean_text(inv.get("reassigned_from")):
        return None
    status = str(inv.get("status") or "").strip().lower()
    if status in _ASSIGNED_STATUSES and inv.get("used_at") not in (None, ""):
        return inv.get("used_at"), "poa_inventory.used_at"
    return None


def _execution_row(
    primary: dict,
    secondary: Optional[dict],
    *,
    raw_date: Any,
    basis: str,
) -> dict:
    sources = [basis]
    if secondary is not None and basis.startswith("poa_inventory"):
        supplemented = False
        for key in (
            "defendant_name", "bond_amount", "amount", "gross_premium",
            "surety_owed", "buf_owed", "case_number", "county", "charge",
            "booking_number",
        ):
            primary_missing = key not in primary or primary.get(key) in (None, "")
            if primary_missing and secondary.get(key) not in (None, ""):
                supplemented = True
                break
        if supplemented:
            sources.append("active_bonds")
    defendant = _stored_defendant(primary) or _stored_defendant(secondary)
    return {
        "poa_number": _norm_poa(_first_stored(primary, secondary, "poa_number", "poa_full")),
        "poa_prefix": _clean_text(_first_stored(primary, secondary, "poa_prefix")),
        "date_executed": _iso_day(raw_date),
        "defendant": defendant,
        "bond_amount": _stored_amount(primary, "bond_amount", "amount")
        if _stored_amount(primary, "bond_amount", "amount") is not None
        else _stored_amount(secondary, "bond_amount", "amount"),
        "gross_premium": _stored_amount(primary, "gross_premium")
        if _stored_amount(primary, "gross_premium") is not None
        else _stored_amount(secondary, "gross_premium"),
        "surety_owed": _stored_amount(primary, "surety_owed")
        if _stored_amount(primary, "surety_owed") is not None
        else _stored_amount(secondary, "surety_owed"),
        "buf_owed": _stored_amount(primary, "buf_owed", "buf")
        if _stored_amount(primary, "buf_owed", "buf") is not None
        else _stored_amount(secondary, "buf_owed", "buf"),
        "status": _clean_text(_first_stored(primary, secondary, "status")),
        "case_number": _clean_text(_first_stored(primary, secondary, "case_number")),
        "booking_number": _clean_text(_first_stored(primary, secondary, "booking_number", "bond_case_id")),
        "county": _clean_text(_first_stored(primary, secondary, "county")),
        "charge": _clean_text(_first_stored(primary, secondary, "charge")),
        "source": "+".join(sources),
    }


def _select_executions(
    inventory: Iterable[dict],
    bonds: Iterable[dict],
    surety_id: str,
    start_dt: Optional[datetime],
    end_dt: Optional[datetime],
) -> tuple[list[dict], int]:
    inv_rows = [d for d in inventory if inventory_surety_matches(d, surety_id)]
    bond_rows = [d for d in bonds if bond_matches_surety(d, surety_id)]
    bonds_by_poa: dict[str, dict] = {}
    for bond in bond_rows:
        poa = _norm_poa(bond.get("poa_number") or bond.get("poa_full"))
        if poa and poa not in bonds_by_poa:
            bonds_by_poa[poa] = bond

    chosen: dict[str, dict] = {}
    omitted_undated = 0
    window = _window_active(start_dt, end_dt)

    for inv in inv_rows:
        poa = _norm_poa(inv.get("poa_number") or inv.get("poa_full"))
        if not poa:
            continue
        stamp = _execution_stamp(inv)
        if stamp is None:
            status = str(inv.get("status") or "").strip().lower()
            if status in _ASSIGNED_STATUSES and window:
                omitted_undated += 1
            elif status in _ASSIGNED_STATUSES and not window:
                chosen[poa] = _execution_row(inv, bonds_by_poa.get(poa), raw_date=None, basis="poa_inventory.status")
            continue
        raw, basis = stamp
        if window and not _in_window(raw, start_dt, end_dt):
            continue
        if window and _iso_day(raw) is None:
            omitted_undated += 1
            continue
        chosen[poa] = _execution_row(inv, bonds_by_poa.get(poa), raw_date=raw, basis=basis)

    for bond in bond_rows:
        poa = _norm_poa(bond.get("poa_number") or bond.get("poa_full"))
        if not poa or poa in chosen:
            continue
        raw = bond.get("bond_date") if bond.get("bond_date") not in (None, "") else bond.get("date_executed")
        if raw in (None, ""):
            if window:
                omitted_undated += 1
            else:
                chosen[poa] = _execution_row(bond, None, raw_date=None, basis="active_bonds.bond_date")
            continue
        if window and not _in_window(raw, start_dt, end_dt):
            continue
        chosen[poa] = _execution_row(bond, None, raw_date=raw, basis="active_bonds.bond_date")

    rows = sorted(chosen.values(), key=lambda r: (r.get("date_executed") or "9999-99-99", r.get("poa_number") or ""))
    return rows, omitted_undated


def _select_voids(
    inventory: Iterable[dict],
    surety_id: str,
    start_dt: Optional[datetime],
    end_dt: Optional[datetime],
) -> tuple[list[dict], int]:
    rows = []
    omitted_undated = 0
    window = _window_active(start_dt, end_dt)
    for doc in inventory:
        if str(doc.get("status") or "") != "voided":
            continue
        if not inventory_surety_matches(doc, surety_id):
            continue
        raw = doc.get("voided_at")
        if window and not _in_window(raw, start_dt, end_dt):
            if _iso_day(raw) is None:
                omitted_undated += 1
            continue
        voided_by = doc.get("voided_by") if "voided_by" in doc and doc.get("voided_by") not in (None, "") else None
        rows.append({
            "poa_number": _norm_poa(doc.get("poa_number") or doc.get("poa_full")),
            "poa_prefix": _clean_text(doc.get("poa_prefix")),
            "surety_id": surety_id,
            "bond_amount": _stored_amount(doc, "bond_amount", "amount"),
            "voided_by": _clean_text(voided_by),
            "void_reason": _clean_text(doc.get("void_reason")),
            "voided_at": _iso_day(raw),
            "source": "poa_inventory.status",
        })
    rows.sort(key=lambda r: (r.get("voided_at") or "9999-99-99", r.get("poa_number") or ""))
    return rows, omitted_undated


def _meaningful_reassign(val: Any) -> bool:
    text = str(val or "").strip()
    return bool(text) and text.lower() not in ("none", "null")


def _release_reason(doc: dict) -> Optional[str]:
    return _clean_text(doc.get("release_reason") or doc.get("released_reason"))


def _renewal_when(bond: dict) -> Any:
    if bond.get("last_renewed_at") not in (None, ""):
        return bond.get("last_renewed_at")
    hist = bond.get("renewal_history")
    if isinstance(hist, list) and hist:
        last = hist[-1]
        if isinstance(last, dict):
            return last.get("renewed_at")
    return None


def _transfer_row(
    *,
    event_date: Optional[str],
    event_type: str,
    poa_number: str,
    from_ref: Optional[str],
    to_ref: Optional[str],
    reason: Optional[str],
    defendant: Optional[str],
    sources: list[str],
) -> dict:
    return {
        "event_date": event_date,
        "event_type": event_type,
        "poa_number": poa_number,
        "from_ref": from_ref,
        "to_ref": to_ref,
        "reason": reason,
        "defendant": defendant,
        "sources": sources,
    }


def _select_transfers(
    inventory: Iterable[dict],
    bonds: Iterable[dict],
    audits: Iterable[dict],
    audit_inventory: Iterable[dict],
    surety_id: str,
    start_dt: Optional[datetime],
    end_dt: Optional[datetime],
) -> tuple[list[dict], int, int]:
    inv_rows = [d for d in inventory if inventory_surety_matches(d, surety_id)]
    audit_by_poa: dict[str, dict] = {}
    for doc in audit_inventory:
        poa = _norm_poa(doc.get("poa_number") or doc.get("poa_full"))
        if poa and poa not in audit_by_poa:
            audit_by_poa[poa] = doc

    omitted_undated = 0
    unattributed = 0
    window = _window_active(start_dt, end_dt)
    rows: list[dict] = []
    release_index: dict[tuple[str, Optional[str]], dict] = {}

    for doc in inv_rows:
        poa = _norm_poa(doc.get("poa_number") or doc.get("poa_full"))
        if not poa:
            continue
        if _meaningful_reassign(doc.get("reassigned_from")):
            raw = doc.get("used_at")
            day = _iso_day(raw)
            if window and not _in_window(raw, start_dt, end_dt):
                if day is None:
                    omitted_undated += 1
            else:
                if day is None and window:
                    omitted_undated += 1
                else:
                    rows.append(_transfer_row(
                        event_date=day,
                        event_type="reassign",
                        poa_number=poa,
                        from_ref=_clean_text(doc.get("reassigned_from")),
                        to_ref=_clean_text(doc.get("bond_case_id")),
                        reason=None,
                        defendant=_stored_defendant(doc),
                        sources=["poa_inventory.reassigned_from"],
                    ))
        if doc.get("released_at") not in (None, ""):
            raw = doc.get("released_at")
            day = _iso_day(raw)
            if window and not _in_window(raw, start_dt, end_dt):
                if day is None:
                    omitted_undated += 1
                continue
            if day is None and window:
                omitted_undated += 1
                continue
            row = _transfer_row(
                event_date=day,
                event_type="release",
                poa_number=poa,
                from_ref=_clean_text(doc.get("bond_case_id")),
                to_ref=None,
                reason=_release_reason(doc),
                defendant=_stored_defendant(doc),
                sources=["poa_inventory.released_at"],
            )
            rows.append(row)
            release_index[(poa, day)] = row

    bond_rows = [d for d in bonds if bond_matches_surety(d, surety_id)]
    for bond in bond_rows:
        prev = _norm_poa(bond.get("previous_poa_number"))
        if not prev:
            continue
        raw = _renewal_when(bond)
        day = _iso_day(raw)
        if window and not _in_window(raw, start_dt, end_dt):
            if day is None:
                omitted_undated += 1
            continue
        if day is None and window:
            omitted_undated += 1
            continue
        current = _norm_poa(bond.get("poa_number"))
        existing = release_index.get((prev, day))
        if existing:
            existing["event_type"] = "renewal_poa_change"
            existing["from_ref"] = prev
            existing["to_ref"] = current or existing.get("to_ref")
            if not existing.get("reason"):
                existing["reason"] = _clean_text(bond.get("renewal_reason"))
            if not existing.get("defendant"):
                existing["defendant"] = _stored_defendant(bond)
            if "active_bonds.previous_poa_number" not in existing["sources"]:
                existing["sources"].append("active_bonds.previous_poa_number")
            continue
        row = _transfer_row(
            event_date=day,
            event_type="renewal_poa_change",
            poa_number=prev,
            from_ref=prev,
            to_ref=current or None,
            reason=_clean_text(bond.get("renewal_reason")),
            defendant=_stored_defendant(bond),
            sources=["active_bonds.previous_poa_number"],
        )
        rows.append(row)
        release_index[(prev, day)] = row

    for event in audits:
        if str(event.get("entity_type") or "") != "poa":
            continue
        action = str(event.get("action") or "")
        if action not in TRANSFER_AUDIT_ACTIONS:
            continue
        poa = _norm_poa(event.get("entity_id"))
        inv = audit_by_poa.get(poa)
        if inv is None:
            # Power is not in inventory, so the carrier cannot be proven.
            unattributed += 1
            continue
        if not inventory_surety_matches(inv, surety_id):
            continue
        raw = event.get("timestamp")
        day = _iso_day(raw)
        if window and not _in_window(raw, start_dt, end_dt):
            if day is None:
                omitted_undated += 1
            continue
        if day is None and window:
            omitted_undated += 1
            continue
        details = event.get("details") if isinstance(event.get("details"), dict) else {}
        existing = release_index.get((poa, day))
        source = f"audit_events.{action}"
        if existing:
            if source not in existing["sources"]:
                existing["sources"].append(source)
            if not existing.get("reason"):
                existing["reason"] = _clean_text(details.get("reason"))
            continue
        from_ref = _clean_text(
            details.get("reassigned_from") or details.get("from") or details.get("old_case")
        )
        to_ref = _clean_text(
            details.get("bond_case_id") or details.get("to") or details.get("new_booking_number")
        )
        row = _transfer_row(
            event_date=day,
            event_type="auto_released" if action == "auto_released" else action,
            poa_number=poa,
            from_ref=from_ref,
            to_ref=to_ref,
            reason=_clean_text(details.get("reason")),
            defendant=_stored_defendant(inv),
            sources=[source],
        )
        rows.append(row)
        if action in ("auto_released", "released"):
            release_index[(poa, day)] = row

    rows.sort(key=lambda r: (
        r.get("event_date") or "9999-99-99",
        r.get("poa_number") or "",
        r.get("event_type") or "",
    ))
    return rows, omitted_undated, unattributed


def _select_liability(
    bonds: Iterable[dict],
    surety_id: str,
    start_dt: Optional[datetime],
    end_dt: Optional[datetime],
) -> tuple[list[dict], int]:
    rows = []
    omitted_undated = 0
    window = _window_active(start_dt, end_dt)
    for bond in bonds:
        if not bond_matches_surety(bond, surety_id):
            continue
        status = str(bond.get("status") or "").strip().lower()
        if status in _CLOSED_BOND_STATUSES:
            continue
        raw = bond.get("bond_date") if bond.get("bond_date") not in (None, "") else bond.get("date_executed")
        if window and not _in_window(raw, start_dt, end_dt):
            if _iso_day(raw) is None:
                omitted_undated += 1
            continue
        rows.append({
            "poa_number": _norm_poa(bond.get("poa_number") or bond.get("poa_full")) or None,
            "date": _iso_day(raw),
            "defendant": _stored_defendant(bond),
            "bond_amount": _stored_amount(bond, "bond_amount", "amount"),
            "status": _clean_text(bond.get("status")),
            "case_number": _clean_text(bond.get("case_number")),
            "county": _clean_text(bond.get("county")),
            "source": "active_bonds",
        })
    rows.sort(key=lambda r: (r.get("date") or "9999-99-99", r.get("poa_number") or ""))
    return rows, omitted_undated


async def assemble_powers_pack(
    db,
    *,
    surety_id: str,
    pack: str,
    start_date: str | None = None,
    end_date: str | None = None,
) -> dict:
    """Build the pack contract. ``surety_id`` must already be an active carrier id."""
    start_dt, end_dt, warnings = parse_report_date_window(start_date, end_date)
    resolved_start = start_dt.strftime("%Y-%m-%d") if start_dt else None
    resolved_end = end_dt.strftime("%Y-%m-%d") if end_dt else None

    inventory, inv_trunc = await _fetch(db, "poa_inventory", _inventory_query(surety_id))
    bond_docs, bond_trunc = await _fetch(db, "active_bonds", _bond_query(surety_id))
    audits, audit_trunc = await _fetch(db, "audit_events", {
        "entity_type": "poa",
        "action": {"$in": list(TRANSFER_AUDIT_ACTIONS)},
    })
    audit_poas = [
        _norm_poa(event.get("entity_id"))
        for event in audits
        if _norm_poa(event.get("entity_id"))
    ]
    if audit_poas:
        audit_inventory, audit_inv_trunc = await _fetch(
            db, "poa_inventory", {"poa_number": {"$in": audit_poas}}
        )
    else:
        audit_inventory, audit_inv_trunc = [], False
    if inv_trunc or bond_trunc or audit_trunc or audit_inv_trunc:
        warnings.append(
            f"query hit row limit ({REPORT_ROW_LIMIT}); narrow the date range if rows look missing"
        )

    executions, exec_undated = _select_executions(inventory, bond_docs, surety_id, start_dt, end_dt)
    voids, void_undated = _select_voids(inventory, surety_id, start_dt, end_dt)
    transfers, xfer_undated, unattributed = _select_transfers(
        inventory, bond_docs, audits, audit_inventory, surety_id, start_dt, end_dt
    )
    liability, liab_undated = _select_liability(bond_docs, surety_id, start_dt, end_dt)

    if exec_undated:
        warnings.append(f"{exec_undated} execution candidate(s) had no date in range and were omitted")
    if void_undated:
        warnings.append(f"{void_undated} voided power(s) had no voided_at and were omitted from the dated range")
    if xfer_undated:
        warnings.append(f"{xfer_undated} transfer evidence row(s) had no event date and were omitted from the dated range")
    if unattributed:
        warnings.append(
            f"{unattributed} poa audit event(s) could not be matched to this carrier's inventory and were omitted"
        )
    if liab_undated:
        warnings.append(f"{liab_undated} open bond(s) had no bond date and were omitted from the dated range")

    on_date = resolved_end or datetime.now(timezone.utc).strftime("%Y-%m-%d")
    carrier = SURETY_REGISTRY[surety_id]["short"]
    contract = {
        "success": True,
        "surety_id": surety_id,
        "carrier": carrier,
        "carrier_label": SURETY_REGISTRY[surety_id]["label"],
        "pack": pack,
        "filename": powers_pack_filename(surety_id, pack, on_date),
        "start_date": resolved_start,
        "end_date": resolved_end,
        "warnings": warnings,
        "history_complete": False,
        "transfer_banner": TRANSFER_HISTORY_BANNER,
        "execution_note": EXECUTION_PACK_NOTE,
        "void_note": VOID_PACK_NOTE,
        "liability_note": LIABILITY_PACK_NOTE,
        "execution": {
            "count": len(executions),
            "rows": executions,
            "omitted_undated": exec_undated,
            "truncated": inv_trunc or bond_trunc,
        },
        "void": {
            "count": len(voids),
            "rows": voids,
            "omitted_undated": void_undated,
            "truncated": inv_trunc,
        },
        "transfer": {
            "count": len(transfers),
            "rows": transfers,
            "omitted_undated": xfer_undated,
            "unattributed_audit_events": unattributed,
            "banner": TRANSFER_HISTORY_BANNER,
            "history_complete": False,
            "truncated": inv_trunc or bond_trunc or audit_trunc or audit_inv_trunc,
        },
        "liability": {
            "count": len(liability),
            "rows": liability,
            "omitted_undated": liab_undated,
            "note": LIABILITY_PACK_NOTE,
            "truncated": bond_trunc,
        },
    }
    logger.info(
        "powers pack surety=%s pack=%s execution=%s void=%s transfer=%s liability=%s",
        surety_id,
        pack,
        len(executions),
        len(voids),
        len(transfers),
        len(liability),
    )
    return contract


def _sheet_for(contract: dict, pack: str) -> list[tuple[str, str, list[str], list[list], str]]:
    """Return (sheet name, note, headers, data rows, empty line) for the requested pack."""
    sid = contract["surety_id"]
    specs = []

    def add(name: str, note: str, headers: list[str], data: list[list], empty: str):
        specs.append((sheet_title(sid, name), note, headers, data, empty))

    if pack in ("execution", "combined"):
        data = []
        for row in contract["execution"]["rows"]:
            data.append([
                row.get("poa_number") or "",
                row.get("poa_prefix") or "",
                row.get("date_executed") or "",
                row.get("defendant") or "",
                row.get("bond_amount"),
                row.get("gross_premium"),
                row.get("surety_owed"),
                row.get("buf_owed"),
                row.get("status") or "",
                row.get("case_number") or "",
                row.get("booking_number") or "",
                row.get("county") or "",
                row.get("charge") or "",
                row.get("source") or "",
            ])
        add(
            "Execution",
            contract["execution_note"],
            [
                "Power #", "Prefix", "Executed", "Defendant", "Bond amount",
                "Gross premium", "Surety owed", "BUF", "Status", "Case number",
                "Booking", "County", "Charge", "Source",
            ],
            data,
            "No powers executed in this range.",
        )
    if pack in ("void", "combined"):
        data = []
        for row in contract["void"]["rows"]:
            data.append([
                row.get("poa_number") or "",
                row.get("surety_id") or "",
                row.get("bond_amount"),
                row.get("voided_by") or "",
                row.get("void_reason") or "",
                row.get("voided_at") or "",
                row.get("source") or "",
            ])
        add(
            "Void",
            contract["void_note"],
            ["Power #", "Surety", "Bond amount", "Voided by", "Reason", "Voided", "Source"],
            data,
            "No voided powers in this range.",
        )
    if pack in ("transfer", "combined"):
        data = []
        for row in contract["transfer"]["rows"]:
            data.append([
                row.get("event_date") or "",
                row.get("event_type") or "",
                row.get("poa_number") or "",
                row.get("from_ref") or "",
                row.get("to_ref") or "",
                row.get("reason") or "",
                row.get("defendant") or "",
                ", ".join(row.get("sources") or []),
            ])
        add(
            "Transfer",
            contract["transfer_banner"],
            ["Event date", "Event type", "Power #", "From", "To", "Reason", "Defendant", "Evidence"],
            data,
            "No transfer evidence in this range.",
        )
    if pack == "combined":
        data = []
        for row in contract["liability"]["rows"]:
            data.append([
                row.get("poa_number") or "",
                row.get("date") or "",
                row.get("defendant") or "",
                row.get("bond_amount"),
                row.get("status") or "",
                row.get("case_number") or "",
                row.get("county") or "",
                row.get("source") or "",
            ])
        add(
            "Liability",
            contract["liability_note"],
            ["Power #", "Bond date", "Defendant", "Bond amount", "Status", "Case number", "County", "Source"],
            data,
            "No open-book bonds in this range.",
        )
    return specs


def build_powers_pack_xlsx(contract: dict, *, pack: str) -> bytes:
    """Multi-sheet workbook for combined; one carrier-named sheet for a single pack."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    sid = contract["surety_id"]
    carrier = contract["carrier"]
    wb = Workbook()
    # Cover is always first so the partial-history banner ships inside the file.
    cover = wb.active
    cover.title = sheet_title(sid, "Cover")
    title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    label_font = Font(name="Calibri", size=11, bold=True, color="0F172A")
    body_font = Font(name="Calibri", size=11, color="0F172A")
    brand = PatternFill("solid", fgColor="0B3D2E")
    cover["A1"] = f"{AGENCY_NAME} — {carrier} powers pack"
    cover["A1"].font = title_font
    cover["A1"].fill = brand
    cover.merge_cells("A1:B1")
    window = "All available records"
    if contract.get("start_date") or contract.get("end_date"):
        window = f"{contract.get('start_date') or '…'} → {contract.get('end_date') or '…'}"
    cover_rows = [
        ("Carrier", contract.get("carrier_label") or carrier),
        ("Pack", pack),
        ("File", contract.get("filename") or ""),
        ("Reporting window", window),
        ("Execution rows", contract["execution"]["count"]),
        ("Void rows", contract["void"]["count"]),
        ("Transfer rows", contract["transfer"]["count"]),
        ("Transfer history complete", "no"),
        ("Liability rows", contract["liability"]["count"] if pack == "combined" else "see Surety Liability report"),
        ("Amounts", "Stored values only. Blank means the field was not on the record."),
        ("Transfer log", contract["transfer_banner"]),
        ("Execution", contract["execution_note"]),
        ("Void", contract["void_note"]),
        ("Liability", contract["liability_note"] if pack == "combined" else "Not included on this single pack."),
    ]
    warnings = contract.get("warnings") or []
    if warnings:
        cover_rows.append(("Warnings", " | ".join(warnings)))
    for i, (label, value) in enumerate(cover_rows, 3):
        cover.cell(i, 1, label).font = label_font
        cell = cover.cell(i, 2, value)
        cell.font = body_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    cover.column_dimensions["A"].width = 28
    cover.column_dimensions["B"].width = 88
    cover.row_dimensions[13].height = 72

    thin = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    header_fill = PatternFill("solid", fgColor="0F172A")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    note_font = Font(name="Calibri", size=9, italic=True, color="334155")
    money_cols_by_header = {"Bond amount", "Gross premium", "Surety owed", "BUF"}

    for title, note, headers, data, empty in _sheet_for(contract, pack):
        ws = wb.create_sheet(title)
        ws["A1"] = f"{carrier} {title.split(' ', 1)[-1]}"
        ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        ws["A1"].fill = brand
        last_col = max(len(headers), 1)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
        note_cell = ws.cell(2, 1, note)
        note_cell.font = note_font
        note_cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = 48
        ws.cell(3, 1, empty if not data else f"Rows: {len(data)}").font = note_font
        header_row = 5
        for col, header in enumerate(headers, 1):
            cell = ws.cell(header_row, col, header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin
        for r_i, values in enumerate(data):
            for c_i, val in enumerate(values, 1):
                cell = ws.cell(header_row + 1 + r_i, c_i, "" if val is None else val)
                cell.border = thin
                cell.font = Font(name="Calibri", size=10)
                header = headers[c_i - 1]
                if header in money_cols_by_header and isinstance(val, (int, float)):
                    cell.number_format = '"$"#,##0.00'
        for i in range(1, last_col + 1):
            ws.column_dimensions[get_column_letter(i)].width = 18
        ws.column_dimensions["A"].width = 22
        ws.freeze_panes = "A6"
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{header_row + max(len(data), 1)}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.oddHeader.left.text = f"{carrier} powers"
        ws.oddFooter.right.text = contract.get("filename") or ""

    buf = io.BytesIO()
    from core.booking_identity import redact_workbook
    redact_workbook(wb)  # XLSX never prints a Miami-Dade internal key
    wb.save(buf)
    return buf.getvalue()
